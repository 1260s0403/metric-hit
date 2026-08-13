from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path("outputs/metric_hit_strategy/MetricHit_strategy_goals_and_keywords.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5D6773"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
GREEN = "1F6B4F"
WHITE = "FFFFFF"
BLACK = "202124"


def set_run(run, size=11, bold=False, color=BLACK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[i]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)


def cell_text(cell, text, bold=False, color=BLACK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, header_fill=LIGHT, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True, color=NAVY, size=font_size)
        shade(table.rows[0].cells[i], header_fill)
    repeat_header(table.rows[0])
    keep_row_together(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell_text(cells[i], val, size=font_size)
            if ridx % 2 == 1:
                shade(cells[i], "FAFBFC")
        keep_row_together(table.rows[-1])
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_para(doc, text="", bold_lead=None, after=6, size=11, color=BLACK, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), size=size, bold=True, color=color)
        set_run(p.add_run(text[len(bold_lead):]), size=size, color=color, italic=italic)
    else:
        set_run(p.add_run(text), size=size, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_number(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(label), bold=True, color=label_color, size=10.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    set_run(p2.add_run(text), size=10.5)
    set_table_geometry(table, [9360])
    keep_row_together(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size={1:16, 2:13, 3:12}[level], bold=True,
            color={1:BLUE, 2:BLUE, 3:DARK_BLUE}[level])
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("Страница "), size=9, color=MUTED)
    begin_run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fldChar1)
    instr_run = paragraph.add_run()
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    instr_run._r.append(instrText)
    sep_run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(fld_sep)
    set_run(paragraph.add_run("1"), size=9, color=MUTED)
    end_run = paragraph.add_run()
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    end_run._r.append(fldChar2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Explicit standard_business_brief styles.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for level, size, before, after, color in [
    (1, 16, 16, 8, BLUE), (2, 13, 12, 6, BLUE), (3, 12, 8, 4, DARK_BLUE)
]:
    st = doc.styles[f"Heading {level}"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Quiet running furniture.
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run(hp.add_run("METRICHIT  |  СТРАТЕГИЯ ПОИСКОВОГО ПРИВЛЕЧЕНИЯ"), size=8.5, bold=True, color=MUTED)
add_page_number(section.footer.paragraphs[0])

# Memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(5)
set_run(p.add_run("СТРАТЕГИЧЕСКИЙ ОТЧЁТ"), size=10, bold=True, color=BLUE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
set_run(p.add_run("Цели и приоритетные поисковые запросы MetricHit"), size=25, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
set_run(p.add_run("Как выбирать ключи по ожидаемой экономической ценности, а не по объёму трафика"), size=13, color=MUTED)

meta = [
    ("Дата", "13 августа 2026"),
    ("Объект", "Поисковое привлечение новых пользователей MetricHit"),
    ("Горизонт", "Первый цикл: 90 дней"),
    ("Статус", "Стратегическая версия; количественная валидация частотности и конверсий обязательна"),
]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(f"{label}: "), bold=True, size=10.5)
    set_run(p.add_run(value), size=10.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
add_callout(
    doc,
    "Главное решение",
    "Считать основной целью не позицию, не трафик и не клик по кнопке, а первый оплаченный баланс нового пользователя, атрибутированный к поисковому запросу. Первые темы продвижения: выбор сервиса, цена, самостоятельный запуск и работа с несколькими проектами.",
    fill=PALE_GOLD,
    label_color="7A5A00",
)

add_heading(doc, "1. Рамка анализа", 1)
add_para(doc, "Отчёт построен без привязки к конкретному сайту публикации. Единицей планирования является не площадка, а поисковая цель: отдельный кластер запросов, единый интент пользователя, одна релевантная страница и одно измеримое следующее действие.")
add_para(doc, "Текущая оценка ключей является стратегической, а не статистическим прогнозом. Доступная семантика позволяет определить коммерческую близость и очередность тестов, но точную прибыльность нельзя честно рассчитать без частотности Wordstat, текущих позиций, стоимости продвижения и фактической конверсии в оплату.")

add_heading(doc, "2. Что именно считать целями", 1)
add_heading(doc, "2.1. Иерархия целей бизнеса", 2)
goal_rows = [
    ("G1", "Первое пополнение", "Новый пользователь внёс деньги", "Главная денежная конверсия"),
    ("G2", "Первый активный проект", "Создан проект и запущена работа", "Показывает реальное использование"),
    ("G3", "Повторное пополнение", "Пользователь вернулся и снова оплатил", "Подтверждает ценность и LTV"),
    ("G4", "Регистрация завершена", "Создан аккаунт", "Основная лид-конверсия, но ещё не выручка"),
    ("G5", "Контакт с поддержкой", "Запрошен бонус или помощь с запуском", "Сильный сигнал намерения"),
    ("G6", "Переход к сервису", "Клик со страницы привлечения", "Микроконверсия для диагностики"),
]
add_table(doc, ["Код", "Цель", "Факт достижения", "Роль"], goal_rows, [650, 2050, 3100, 3560], font_size=8.8)

add_callout(doc, "North Star", "Валовая маржа от новых пользователей, привлечённых из поиска. Пока маржа и повторные платежи не связаны с источником, рабочий заменитель — сумма первых пополнений минус затраты на создание и продвижение страниц.")

add_heading(doc, "2.2. Что целью не является", 2)
add_bullet(doc, "Позиция страницы сама по себе: она лишь создаёт возможность получить показы.")
add_bullet(doc, "Трафик: большой поток по информационному запросу может не давать регистраций.")
add_bullet(doc, "CTR сниппета: высокий CTR бесполезен, если ожидание не совпало с содержанием.")
add_bullet(doc, "Клик «Регистрация»: это намерение, а не завершённая регистрация.")
add_bullet(doc, "Бесплатные клики: активация бонуса полезна для онбординга, но не доказывает платёжеспособность.")

add_heading(doc, "2.3. Воронка измерения", 2)
funnel_rows = [
    ("1", "Показ в поиске", "Запрос, позиция, URL", "Видимость"),
    ("2", "Переход на страницу", "CTR и органический визит", "Соответствие сниппета"),
    ("3", "Переход в MetricHit", "Клик по размеченной ссылке", "Интерес к продукту"),
    ("4", "Регистрация завершена", "Событие после создания аккаунта", "Лид"),
    ("5", "Обращение/активация", "Контакт с поддержкой", "Готовность начать"),
    ("6", "Первое пополнение", "Сумма и идентификатор пользователя", "Выручка"),
    ("7", "Первый запуск", "Проект создан и активирован", "Активация продукта"),
    ("8", "Повторное пополнение", "Сумма в окне 30/60/90 дней", "Удержание и LTV"),
]
add_table(doc, ["Этап", "Событие", "Что фиксировать", "Смысл"], funnel_rows, [650, 2350, 3500, 2860], font_size=8.6)

p = add_heading(doc, "3. Кому и какую задачу продаёт MetricHit", 1)
p.paragraph_format.page_break_before = True
audience_rows = [
    ("SEO-специалист", "Управлять несколькими проектами без посредника", "Сервис, автоматизация, несколько сайтов", "Высокая"),
    ("Агентство", "Масштабировать работу и контролировать бюджет", "Для агентств, для нескольких проектов", "Очень высокая LTV"),
    ("Владелец сайта", "Сдвинуть застрявшие позиции", "Сайт не растёт, вторая страница, ТОП-20", "Средняя"),
    ("Опытный SEO-заказчик", "Сравнить решения и стоимость", "Цена, стоимость, цена за клик", "Высокая"),
    ("Новичок", "Понять принцип и риски", "Что такое, как работает, риски", "Низкая сразу; прогрев"),
]
add_table(doc, ["Сегмент", "Задача", "Язык запроса", "Ценность"], audience_rows, [1600, 2850, 3100, 1810], font_size=8.6)

add_heading(doc, "4. Как определять выгодность ключа", 1)
add_para(doc, "Выгодность — это ожидаемая прибыль на единицу затрат, а не максимальная частотность. Для каждого кластера применяется двухступенчатая оценка.")
add_heading(doc, "4.1. Стратегический балл до запуска", 2)
score_rows = [
    ("Коммерческий интент", "25%", "Насколько запрос близок к выбору и покупке"),
    ("Соответствие продукту", "20%", "Есть ли прямой и честный ответ в возможностях MetricHit"),
    ("Потенциал конверсии", "15%", "Вероятность регистрации и первого платежа"),
    ("Потенциал спроса", "15%", "Частотность, сезонность и число вариаций"),
    ("Ценность клиента", "10%", "Вероятный чек, число проектов, повторные пополнения"),
    ("Достижимость выдачи", "10%", "Сила конкурентов и соответствие формату страницы"),
    ("Контентная ясность", "5%", "Можно ли дать полный ответ без размывания интента"),
    ("Штраф за риск", "0–15 п.", "Политики, чрезмерные обещания, нестабильность или нерелевантность"),
]
add_table(doc, ["Фактор", "Вес", "Что оцениваем"], score_rows, [2600, 950, 5810], font_size=9)
add_para(doc, "Формула первого этапа: сумма взвешенных оценок по шкале 0–5, нормированная до 100, минус штраф за риск. Это приоритизация экспериментов, а не обещание результата.", italic=True, color=MUTED, size=10)

p = add_heading(doc, "4.2. Экономический балл после запуска", 2)
p.paragraph_format.page_break_before = True
add_callout(doc, "Рабочая формула", "Экономическая ценность кластера = (поисковые визиты × CR в регистрацию × CR в первое пополнение × средняя маржа на 90 дней × вероятность удержания) − полная стоимость создания и продвижения.", fill=PALE_GOLD, label_color="7A5A00")
add_para(doc, "Через 30–90 дней стратегический балл должен уступить место фактической экономике. Кластер с меньшим трафиком, но высокой долей оплат, получает больше бюджета, чем массовый информационный кластер.")

add_heading(doc, "5. Приоритетные поисковые цели", 1)
priority_rows = [
    ("1", "Цена и бюджет", "87", "Очень высокий", "Лучший баланс намерения, ясности оффера и измеримости"),
    ("2", "Выбор сервиса", "85", "Очень высокий", "Пользователь уже выбирает инструмент"),
    ("3", "Самостоятельный запуск", "79", "Высокий", "Прямой доступ к сервису соответствует задаче"),
    ("4", "Агентства и несколько проектов", "78", "Высокий", "Вероятно меньший спрос, но выше LTV"),
    ("5", "Основная категория ПФ", "75", "Высокий", "Большой потенциал, но конкуренция и риск выше"),
    ("6", "Риски и последствия", "64", "Средний", "Сильный этап принятия решения; нужен осторожный тон"),
    ("7", "Застрявшие позиции", "60", "Средний", "Широкая проблема, но не всегда ведёт к продукту"),
    ("8", "Брендовые запросы", "58*", "Очень высокий", "Низкий объём привлечения; обязательная защита спроса"),
    ("9", "Обучающие запросы", "47", "Низкий сразу", "Полезны для охвата и внутреннего прогрева"),
]
add_table(doc, ["№", "Поисковая цель", "Балл", "Коммерческий потенциал", "Вывод"], priority_rows, [450, 2050, 650, 1900, 4310], font_size=8.2)
add_para(doc, "* Брендовый кластер имеет невысокий балл как источник нового небрендового спроса, но максимальный приоритет по контролю репутации и конверсии уже сформированного спроса.", italic=True, color=MUTED, size=9.5)

p = add_heading(doc, "6. Ядро первой очереди: самые выгодные ключи", 1)
p.paragraph_format.page_break_before = True
p.paragraph_format.space_before = Pt(36)
add_para(doc, "Продвигать следует кластеры, а не отдельные словоформы. Один кластер получает одну основную страницу, один ведущий запрос и набор близких формулировок.")

clusters = [
    ("A1. Выбор сервиса", "сервис накрутки поведенческих факторов", "сервис накрутки ПФ; сервис поведенческих факторов; выбрать сервис ПФ", "Прямой доступ, управление, оплата за клик, контроль проекта", "Регистрация"),
    ("A2. Цена", "накрутка ПФ цена", "стоимость накрутки ПФ; цена клика ПФ; накрутка ПФ цена за клик; бюджет на ПФ", "Тарифная сетка, примеры расчёта, что влияет на расход", "Расчёт → регистрация"),
    ("A3. Самостоятельный запуск", "накрутка ПФ самостоятельно", "как запустить накрутку ПФ; сервис для запуска ПФ; настроить проект ПФ", "Пошаговый запуск через интерфейс без посредника", "Регистрация"),
    ("A4. Профессионалы", "сервис ПФ для SEO-специалистов", "сервис ПФ для агентств; сервис ПФ для нескольких сайтов; несколько проектов ПФ", "Масштаб, контроль расходов, индивидуальные условия на объёме", "Регистрация / контакт"),
    ("A5. Основная категория", "накрутка ПФ Яндекс", "накрутка поведенческих факторов Яндекс; накрутка ПФ сайта; продвижение поведенческими факторами", "Полное объяснение метода, ограничений, рисков и сценария применения", "Регистрация после прогрева"),
]
add_table(doc, ["Кластер", "Ведущий ключ", "Поддерживающие ключи", "Коммерческий ответ", "CTA"], clusters, [1350, 1900, 2650, 2550, 910], font_size=7.9)

add_heading(doc, "7. Вторая очередь: ключи для прогрева и снятия возражений", 1)
second_rows = [
    ("Риски", "накрутка ПФ риски", "последствия накрутки ПФ; санкции за ПФ; можно ли использовать ПФ", "Не обещать безопасность; объяснять условия и ограничения"),
    ("Диагностика", "сайт застрял на второй странице Яндекса", "сайт застрял в ТОП-20; позиции не растут; SEO не даёт роста", "Сначала исключить релевантность, технику и слабый оффер"),
    ("Подбор запросов", "как выбрать запросы для накрутки ПФ", "какие запросы продвигать; частотность для ПФ; распределение запросов", "Мост к созданию проекта и бюджету"),
    ("Управление", "как оценить результат накрутки ПФ", "через сколько растут позиции; когда остановить ПФ; динамика позиций", "Показать дисциплину измерения без гарантий"),
    ("Бренд", "MetricHit", "Metric Hit; Метрик Хит; mtrhit; MetricHit отзывы; MetricHit цена", "Факты, тарифы, инструкция, контакты; без вымышленных отзывов"),
]
add_table(doc, ["Кластер", "Ведущий ключ", "Расширение", "Роль контента"], second_rows, [1200, 2250, 3050, 2860], font_size=8.1)

p = add_heading(doc, "8. Какие запросы пока не брать в продвижение", 1)
p.paragraph_format.page_break_before = True
p.paragraph_format.space_before = Pt(36)
defer_rows = [
    ("что такое поведенческие факторы", "Большой информационный охват, слабое намерение купить", "Оставить поддерживающим материалом"),
    ("как улучшить ПФ без накрутки", "Интент не совпадает с основным продуктом", "Использовать для доверия, не как денежную цель"),
    ("как вывести сайт в ТОП Яндекса", "Слишком широкий и конкурентный запрос", "Дробить на конкретные проблемы"),
    ("накрутка ПФ без риска / безопасная накрутка", "Формирует ожидание, которое нельзя честно гарантировать", "Отвечать через риски, не обещание"),
    ("лучший сервис ПФ", "Нужна доказательная сравнительная база", "Вернуться после накопления кейсов и критериев"),
    ("как выйти из-под фильтра за ПФ", "Слабое соответствие текущему продукту и высокая чувствительность", "Не приоритет"),
]
add_table(doc, ["Запрос", "Почему не сейчас", "Решение"], defer_rows, [2650, 3950, 2760], font_size=8.6)

p = add_heading(doc, "9. Архитектура семантики и защита от каннибализации", 1)
p.paragraph_format.page_break_before = True
p.paragraph_format.space_before = Pt(36)
add_number(doc, "Одна поисковая цель — одна основная страница. Синонимы и близкие формулировки объединяются, а не получают отдельные почти одинаковые материалы.", bold_lead="Одна поисковая цель — одна основная страница.")
add_number(doc, "Разные намерения разделяются. «Цена», «выбор сервиса», «риски» и «самостоятельный запуск» требуют разных ответов и разных CTA.", bold_lead="Разные намерения разделяются.")
add_number(doc, "Каждая новая тема проверяется на пересечение. Если совпадает более половины ключей и поисковая задача одна, новый материал не нужен — обновляется существующий.", bold_lead="Каждая новая тема проверяется на пересечение.")
add_number(doc, "Коммерческая страница получает факты продукта. Информационная — объяснение и переход к следующему шагу; нельзя превращать каждую статью в одинаковую рекламу.", bold_lead="Коммерческая страница получает факты продукта.")
add_number(doc, "Брендовые варианты ведут в единый официальный контур. Название, домены, Telegram и тарифы должны быть одинаковыми во всех материалах.", bold_lead="Брендовые варианты ведут в единый официальный контур.")

add_heading(doc, "10. План работы на 90 дней", 1)
plan_rows = [
    ("Недели 1–2", "Измерение", "Связать запрос/URL/UTM с регистрацией, первым платежом и первым запуском; снять Wordstat и текущие позиции", "Есть базовая воронка и нулевые значения"),
    ("Недели 2–4", "Первая очередь", "Запустить кластеры «выбор сервиса» и «цена»; подготовить разные офферы и CTA", "2 независимых денежных теста"),
    ("Недели 4–6", "Сегментация", "Запустить «самостоятельный старт» и «для специалистов/агентств»", "Проверка CR и LTV разных аудиторий"),
    ("Недели 6–8", "Расширение", "Добавить основную категорию и запросы о рисках", "Рост охвата без потери точности"),
    ("Недели 8–12", "Перераспределение", "Сравнить стоимость первого платежа и 90-дневную ценность; усилить победителей, остановить слабые связки", "Бюджет идёт в прибыльные кластеры"),
]
add_table(doc, ["Период", "Фокус", "Действия", "Результат"], plan_rows, [1150, 1500, 4400, 2310], font_size=8.2)

add_heading(doc, "11. Система аналитики", 1)
add_heading(doc, "11.1. Обязательные события", 2)
events = [
    ("article_to_landing", "Переход со страницы привлечения", "source_url, query_cluster, content_id"),
    ("registration_start", "Начало регистрации", "source_url, cluster"),
    ("registration_complete", "Аккаунт создан", "user_id, cluster"),
    ("support_contact", "Переход/обращение в поддержку", "user_id, reason"),
    ("bonus_activated", "Начислен тестовый пакет", "user_id"),
    ("first_topup", "Первое пополнение", "user_id, amount, currency, cluster"),
    ("first_project_launch", "Первый проект запущен", "user_id, project_type"),
    ("repeat_topup", "Повторное пополнение", "user_id, amount, days_since_first"),
]
add_table(doc, ["Событие", "Что означает", "Минимальные параметры"], events, [2200, 3100, 4060], font_size=8.4)
add_para(doc, "Клики регистрации, входа и поддержки полезны как промежуточные сигналы, но их необходимо отделять от завершённой регистрации и оплаты. Яндекс Метрика поддерживает отдельные целевые события, параметры и передачу дохода; проверка реализации обязательна до оценки ключей.", size=10.5)

add_heading(doc, "11.2. Отчёт по каждому кластеру", 2)
for item in [
    "Частотность и сезонность ведущего ключа и вариантов.",
    "Позиции и доля показов по каждому URL.",
    "CTR выдачи и доля переходов в продукт.",
    "CR: визит → завершённая регистрация → первое пополнение → первый запуск.",
    "Среднее первое пополнение и сумма пополнений за 30/60/90 дней.",
    "Полная стоимость кластера и стоимость первого платящего пользователя.",
    "Причины остановки: нет видимости, нет кликов, нет продуктовых переходов или нет оплат.",
]:
    add_bullet(doc, item)

add_heading(doc, "12. Правила принятия решений", 1)
decision_rows = [
    ("Нет роста показов/позиций", "Проблема видимости или соответствия выдаче", "Перепроверить интент, конкурентность, индексацию и формат"),
    ("Позиции есть, CTR низкий", "Слабый сниппет или неверное обещание", "Изменить заголовок/описание, не трогая интент"),
    ("CTR хороший, переходов в продукт мало", "Страница не создаёт коммерческого моста", "Уточнить оффер, доказательства и CTA"),
    ("Регистрации есть, оплат нет", "Неверная аудитория или проблема онбординга/ценности", "Анализировать сегмент и первые действия в продукте"),
    ("Оплаты есть, повторов нет", "Проблема результата, поддержки или экономики клиента", "Изучить активацию и причины прекращения"),
    ("Мало трафика, но высокий LTV", "Ценный узкий кластер", "Удерживать и расширять близкими long-tail запросами"),
]
add_table(doc, ["Сигнал", "Интерпретация", "Действие"], decision_rows, [2350, 3250, 3760], font_size=8.4)

p = add_heading(doc, "13. Риски и ограничения", 1)
p.paragraph_format.page_break_before = True
p.paragraph_format.space_before = Pt(72)
add_bullet(doc, "Искусственная имитация пользовательских действий может рассматриваться поисковой системой как нарушение и привести к ограничению ранжирования. Нельзя обещать «без риска» или гарантированный ТОП.")
add_bullet(doc, "Прибыльность ключа зависит не только от позиции: важны сниппет, соответствие страницы, доверие, онбординг, тариф и качество продукта.")
add_bullet(doc, "Частотность без очистки может завышать потенциал; нужны регион, тип соответствия, сезонность и исключение нерелевантных формулировок.")
add_bullet(doc, "Внешние страницы и поисковая выдача нестабильны; актив нельзя считать полностью контролируемым.")
add_bullet(doc, "Недостаточно измерять одну сессию. Для агентств и специалистов ценность проявляется в повторных пополнениях и нескольких проектах.")

p = add_heading(doc, "14. Итоговое решение", 1)
add_callout(doc, "Первая волна", "1) «сервис накрутки поведенческих факторов»; 2) «накрутка ПФ цена»; 3) «накрутка ПФ самостоятельно»; 4) «сервис ПФ для SEO-специалистов / агентств»; 5) «накрутка ПФ Яндекс». Каждый пункт — отдельная поисковая цель и отдельный эксперимент.", fill=PALE_GOLD, label_color="7A5A00")
add_para(doc, "Сначала запускаются кластеры, в которых пользователь уже выбирает инструмент или считает бюджет. Затем — сегментные запросы с высоким потенциальным LTV. Основная категория даёт больший охват, но идёт позже из-за конкуренции и более смешанного интента. Информационные запросы используются для прогрева, перелинковки и снятия возражений, а не как главный источник бюджета.")
add_para(doc, "Окончательный список «самых выгодных» ключей фиксируется только после добавления трёх групп данных: точной частотности, фактической стоимости достижения видимости и конверсии до первого/повторного платежа. До этого корректное решение — не масштабирование, а последовательные тесты кластеров первой волны.")

add_heading(doc, "Приложение A. Семантическое ядро первого цикла", 1)
appendix_rows = [
    ("Выбор сервиса", "сервис накрутки поведенческих факторов; сервис накрутки ПФ; сервис поведенческих факторов; выбрать сервис ПФ; заказать накрутку ПФ"),
    ("Цена", "накрутка ПФ цена; стоимость накрутки ПФ; цена клика ПФ; накрутка ПФ цена за клик; бюджет на накрутку ПФ; как рассчитать бюджет ПФ"),
    ("Самостоятельный запуск", "накрутка ПФ самостоятельно; как запустить накрутку ПФ; как настроить проект ПФ; сервис для запуска ПФ; автоматизация ПФ"),
    ("Профессиональные сегменты", "сервис ПФ для SEO-специалистов; сервис ПФ для агентств; сервис ПФ для нескольких сайтов; продвижение нескольких проектов ПФ"),
    ("Основная категория", "накрутка ПФ Яндекс; накрутка поведенческих факторов Яндекс; накрутка ПФ сайта; накрутка поведенческих факторов сайта; продвижение сайта поведенческими факторами"),
    ("Риски", "накрутка ПФ риски; последствия накрутки ПФ; санкции за накрутку ПФ; можно ли использовать ПФ"),
    ("Диагностика", "сайт застрял на второй странице Яндекса; сайт застрял в ТОП-20; позиции сайта не растут; SEO-трафик есть, заявок нет"),
    ("Подбор и управление", "как выбрать запросы для накрутки ПФ; какие запросы продвигать; как оценить результат ПФ; когда остановить ПФ; через сколько меняются позиции"),
    ("Бренд", "MetricHit; Metric Hit; Метрик Хит; mtrhit; MetricHit цена; MetricHit отзывы"),
]
add_table(doc, ["Кластер", "Ключи"], appendix_rows, [2200, 7160], font_size=8.4)

p = add_heading(doc, "Источники и данные для количественной валидации", 1)
sources = [
    "Яндекс Wordstat — частотность, регионы и сезонность: https://wordstat.yandex.ru/",
    "Яндекс Метрика — цели и типы целей: https://yandex.ru/support/metrica/ru/general/goals",
    "Яндекс Метрика — целевые события и передача дохода: https://yandex.ru/support/metrica/ru/general/goal-js-event",
    "Яндекс Метрика — проверка целей: https://yandex.ru/support/metrica/ru/general/check-goal",
    "Внутренние данные MetricHit — регистрации, пополнения, проекты, повторные платежи и стоимость продвижения.",
]
for s in sources:
    add_bullet(doc, s)

add_para(doc, "Примечание: документ не содержит гарантий позиций или безопасности метода. Оценки 0–100 — экспертная приоритизация для постановки экспериментов; после накопления данных они заменяются фактической экономикой.", italic=True, color=MUTED, size=9.5, after=0)

doc.core_properties.title = "Цели и приоритетные поисковые запросы MetricHit"
doc.core_properties.subject = "Стратегия поискового привлечения"
doc.core_properties.author = "MetricHit"
doc.core_properties.keywords = "MetricHit, SEO, семантика, ключевые запросы, цели, конверсия"
doc.save(OUT)
print(OUT.resolve())
